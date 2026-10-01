from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "documents/resume/Htet-Ko-Oo-Resume.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_cell_margin(cell, top=70, start=110, bottom=70, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=9.4, bold=False, color="142735"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=9.4, bold=False, color="142735"):
    run = paragraph.add_run(text)
    set_font(run, size, bold, color)
    return run


def add_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "235BCC")
    borders.append(bottom)
    p_pr.append(borders)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text.upper(), 9.2, True, "235BCC")
    add_rule(p)


def project(doc, title, tech, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    add_text(p, title, 10, True)
    add_text(p, "  |  " + tech, 8.7, False, "526473")
    for bullet in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(1)
        add_text(p, bullet, 8.8, False)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.43)
section.bottom_margin = Inches(0.43)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(9.4)

header = doc.add_paragraph()
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header.paragraph_format.space_after = Pt(1)
add_text(header, "HTET KO OO", 21, True, "142735")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(4)
add_text(subtitle, "Web Developer Internship Candidate", 10.2, True, "235BCC")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(7)
add_text(contact, "Bangkok, Thailand  |  htetkooo2532@gmail.com  |  github.com/HtetKoOo  |  linkedin.com/in/htet-ko-oo-602913315", 8.5, False, "526473")

section_heading(doc, "Profile")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
add_text(p, "Third-year Digital Technology Innovation student at Kasem Bundit University seeking a web development internship. I build responsive web applications with Next.js, TypeScript, React, and PostgreSQL, with experience in authentication, role-based workflows, data access controls, and practical user interfaces.", 9.2)

section_heading(doc, "Technical Skills")
table = doc.add_table(rows=2, cols=2)
table.autofit = False
table.columns[0].width = Inches(3.55)
table.columns[1].width = Inches(3.55)
skill_rows = [
    ("Frontend", "Next.js, React, TypeScript, JavaScript, HTML, CSS, Tailwind CSS"),
    ("Backend and data", "PostgreSQL, Supabase, Prisma, Drizzle ORM, REST APIs, Laravel"),
    ("Workflow", "Git, GitHub, pnpm, Vercel, Postman, Zod"),
    ("Additional", "Better Auth, Row Level Security, browser APIs, responsive design"),
]
for index, (label, value) in enumerate(skill_rows):
    cell = table.cell(index // 2, index % 2)
    shade(cell, "F3F6F8")
    set_cell_margin(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, label + ": ", 8.5, True, "142735")
    add_text(p, value, 8.5, False, "526473")

section_heading(doc, "Selected Projects")
project(doc, "DayFlow", "Next.js, TypeScript, Supabase, PostgreSQL", [
    "Built a responsive daily-planning PWA with authenticated task workflows, recurring routines, drag-and-drop scheduling, and Day, 2-day, and Week views.",
    "Implemented Supabase Auth, Row Level Security, realtime updates, Zod validation, unit tests, and database isolation checks.",
])
project(doc, "KBU Smart Attendance System", "Next.js, TypeScript, Prisma, PostgreSQL", [
    "Developed role-based admin, lecturer, and student workflows for attendance, schedules, courses, enrollments, and reports.",
    "Implemented browser-local face-template matching with server-side authorization, schedule and enrollment validation, and duplicate-safe attendance records.",
])
project(doc, "Our Sweet Universe", "Next.js, Better Auth, Drizzle ORM, Neon PostgreSQL", [
    "Built a privacy-focused two-person web app with couple-scoped authorization, validated private-memory workflows, and versioned database migrations.",
    "Designed input validation and private media handling foundations for shared settings and personal content.",
])

section_heading(doc, "Education")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
add_text(p, "Kasem Bundit University", 9.5, True)
add_text(p, "  |  Bangkok, Thailand", 8.8, False, "526473")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
add_text(p, "Bachelor's studies in Digital Technology Innovation - Third year, currently enrolled", 9.0, False, "526473")

doc.core_properties.title = "Htet Ko Oo Resume"
doc.core_properties.author = "Htet Ko Oo"
doc.core_properties.subject = "Web developer internship resume"
doc.save(OUT)
